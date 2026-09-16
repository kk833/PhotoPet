"""把 Excel / Word / 文本里的日程计划读进来 (v0.17)。

用户的原话是"我 excel 或者 word 直接丢给他自动生成日程提醒"。这里分两层:

1. **读文件** (`read_file`): .xlsx/.xls/.docx/.doc/.csv/.txt -> 保留行结构的纯文本。
   表格保持"一行一件事"的形状 —— 后面不管是规则解析还是喂给大模型, 都不该丢结构。

2. **解析** (`parse_by_rules` + 交给大模型):
   - **整齐的表格用规则直接解析**: 有表头(日期/时间/事项)的直接按列取, 零成本、
     离线可用、毫秒级。绝大多数人做的日程表就是这种。
   - **规则啃不动的交给大模型**: "下周三下午三点开会"这种自然语言、或者 Word 里
     混在正文里的安排。大模型那条路**复用已有的工具调用 schema**, 只是把写入目标
     换成临时文件 —— 这样它调 add_event 时收集到的是"待确认清单"而不是直接落库。

**读文件不需要任何新依赖**: openpyxl / python-docx / xlrd 都已经在环境里了。
老的 .doc/.xls 走 win32com(本机 Office), 没装 Office 就提示另存为新格式。
"""
import csv
import io
import os
import re

# 表头里出现这些词, 就认为那一列是什么
COL_HINTS = {
    "date": ("日期", "时间点", "date", "日子", "哪天"),
    "time": ("时间", "几点", "时刻", "time", "钟点"),
    "title": ("事项", "事情", "内容", "安排", "任务", "标题", "活动", "会议",
              "title", "event", "subject"),
    "repeat": ("重复", "周期", "频率", "repeat"),
    "note": ("备注", "说明", "note", "备注信息"),
}
MAX_TEXT_CHARS = 12000          # 喂给模型的文本上限, 防止一个大表把 token 烧穿


class ReadResult:
    """读出来的东西 + 它是怎么来的 (给预览窗和报错用)。"""

    def __init__(self, text: str, kind: str, note: str = ""):
        self.text = text            # 保留结构的纯文本
        self.kind = kind            # table / text
        self.note = note            # 读的时候遇到的问题(比如被截断)

    @property
    def empty(self) -> bool:
        return not self.text.strip()


def read_file(path: str) -> ReadResult:
    """把文件读成纯文本。认不出/读不了就抛 ValueError, 由调用方提示用户。"""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _read_excel(path)
    if ext == ".xls":
        return _read_excel_old(path)
    if ext in (".docx",):
        return _read_docx(path)
    if ext in (".doc",):
        return _read_doc_old(path)
    if ext in (".csv", ".txt", ".md"):
        return _read_plain(path)
    raise ValueError(f"这个格式我还不认识：{ext or '（没有扩展名）'}"
                     f"。支持 .xlsx / .xls / .docx / .doc / .csv / .txt")


# ---------- 各种格式 ----------
def _read_excel(path: str) -> ReadResult:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    blocks, truncated = [], False
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c).strip() for c in row]
            if not any(cells):
                continue
            rows.append(" | ".join(cells).rstrip(" |"))
        if rows:
            blocks.append(f"【工作表：{ws.title}】\n" + "\n".join(rows))
    wb.close()
    text = "\n\n".join(blocks)
    if len(text) > MAX_TEXT_CHARS:
        text, truncated = text[:MAX_TEXT_CHARS], True
    note = "表格太长，只读了前面一部分" if truncated else ""
    return ReadResult(text, "table", note)


def _read_excel_old(path: str) -> ReadResult:
    """老的 .xls。先试 xlrd(不用装 Office), 不行再走 win32com。"""
    try:
        import xlrd
        book = xlrd.open_workbook(path)
        blocks = []
        for sheet in book.sheets():
            rows = []
            for r in range(sheet.nrows):
                cells = [str(c).strip() for c in sheet.row_values(r)]
                if any(cells):
                    rows.append(" | ".join(cells).rstrip(" |"))
            if rows:
                blocks.append(f"【工作表：{sheet.name}】\n" + "\n".join(rows))
        text = "\n\n".join(blocks)[:MAX_TEXT_CHARS]
        return ReadResult(text, "table")
    except Exception:                                     # noqa: BLE001
        return _read_via_office(path, "Excel.Application")


def _read_docx(path: str) -> ReadResult:
    import docx
    d = docx.Document(path)
    parts = []
    for para in d.paragraphs:
        line = para.text.strip()
        if line:
            parts.append(line)
    for i, table in enumerate(d.tables, 1):
        rows = []
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells).rstrip(" |"))
        if rows:
            parts.append(f"【表格 {i}】\n" + "\n".join(rows))
    text = "\n".join(parts)
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS]
        return ReadResult(text, "text", "文档太长，只读了前面一部分")
    return ReadResult(text, "text")


def _read_doc_old(path: str) -> ReadResult:
    """老的 .doc。只能借本机 Word (win32com)。"""
    return _read_via_office(path, "Word.Application")


def _read_via_office(path: str, app_name: str) -> ReadResult:
    """用本机 Office 打开另存为纯文本 —— 没装 Office 就明确报错, 别让用户猜。"""
    try:
        import win32com.client as win32
    except ImportError as e:
        raise ValueError("读老格式 .doc/.xls 需要本机装 Office；"
                         "或者用 Word/Excel 另存为 .docx/.xlsx 再丢给我") from e
    abs_path = os.path.abspath(path)
    ext = os.path.splitext(path)[1].lower()
    try:
        app = win32.Dispatch(app_name)
        app.Visible = False
        if ext in (".doc", ".docm", ".docx"):
            doc = app.Documents.Open(abs_path, ReadOnly=True)
            try:
                text = doc.Content.Text
            finally:
                doc.Close(False)
        else:
            wb = app.Workbooks.Open(abs_path, ReadOnly=True)
            try:
                chunks = []
                for sheet in wb.Worksheets:
                    used = sheet.UsedRange
                    for row in used.Rows:
                        cells = [str(c.Text).strip() for c in row.Cells]
                        if any(cells):
                            chunks.append(" | ".join(cells).rstrip(" |"))
                text = "\n".join(chunks)
            finally:
                wb.Close(False)
        app.Quit()
    except Exception as e:                                # noqa: BLE001
        raise ValueError(f"打不开这个文件（{e.__class__.__name__}）。"
                         f"如果本机没装 Office，用 Word/Excel 另存为 .docx/.xlsx 再试") from e
    text = re.sub(r"[\r\x07]", "\n", text)[:MAX_TEXT_CHARS]
    return ReadResult(text, "text" if ext in (".doc", ".docm", ".docx") else "table")


def _read_plain(path: str) -> ReadResult:
    for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
        try:
            with open(path, encoding=enc) as f:
                text = f.read()
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("这个文本文件的编码我认不出来，另存成 UTF-8 再试")
    kind = "table" if path.lower().endswith(".csv") else "text"
    if kind == "table":
        # CSV 走真正的 csv 解析再拼成 " | " —— 直接拿逗号切会把中文句子里
        # 的逗号也切开 (而且 csv 里本来就有引号转义、逗号在字段内这些情况)
        rows = []
        for row in csv.reader(io.StringIO(text)):
            cells = [c.strip() for c in row]
            if any(cells):
                rows.append(" | ".join(cells).rstrip(" |"))
        text = "\n".join(rows)
    if len(text) > MAX_TEXT_CHARS:
        return ReadResult(text[:MAX_TEXT_CHARS], kind, "文件太长，只读了前面一部分")
    return ReadResult(text, kind)


# ---------- 规则解析: 表格 ----------
DATE_RE = re.compile(r"(\d{4})[-/年.](\d{1,2})[-/月.](\d{1,2})")
MD_RE = re.compile(r"(?<!\d)(\d{1,2})[-/月.](\d{1,2})(?!\d)")
# 分钟允许一位数: 手写的表里 "9:5" 很常见, 按两位数字严格匹配会整条丢掉
TIME_RE = re.compile(r"(?<!\d)(\d{1,2})[:：点时](\d{1,2})?(?!\d)")


def _find_col(header: list[str], key: str) -> int | None:
    for i, cell in enumerate(header):
        low = cell.lower()
        for hint in COL_HINTS[key]:
            if hint in low:
                return i
    return None


def parse_by_rules(text: str) -> list[dict]:
    """能按表格规矩解析出来就返回事件列表, 解析不了返回 []（交给大模型）。

    只认**有表头**的表: 没有表头就靠猜列, 猜错的代价是"记了一堆错日程",
    那还不如干脆交给大模型。
    返回 [{"title":…, "date":…, "time":…, "repeat":…}, …]
    """
    events, header, cols = [], None, {}
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line or line.startswith("【"):
            continue
        cells = [c.strip() for c in line.split("|")] if "|" in line else \
            [c.strip() for c in re.split(r"\t+|\s{2,}", line)]
        if len(cells) < 2:
            continue
        if header is None:
            # 五列都要登记 —— 只记 date/time/title 的话, "重复"和"备注"两列
            # 会被静默丢掉 (踩过: 明明写了"每天", 导进来全变成只一次)
            found = {k: _find_col(cells, k) for k in COL_HINTS}
            if found["title"] is not None and (found["date"] is not None
                                               or found["time"] is not None):
                header, cols = cells, found
            continue
        def cell(key):
            idx = cols.get(key)
            return cells[idx] if idx is not None and idx < len(cells) else ""
        title = cell("title")
        date_text, time_text = cell("date"), cell("time")
        if not title or not (date_text or time_text):
            continue
        date_iso = normalize_date(date_text)
        time_iso = normalize_time(time_text)
        if not date_iso:
            continue
        events.append({
            "title": title,
            "date": date_iso,
            "time": time_iso,
            "repeat": normalize_repeat(cell("repeat")),
            "note": cell("note"),
        })
    return events


def normalize_date(text: str, today=None) -> str:
    """把各种写法归一成 YYYY-MM-DD; 认不出返回空串。"""
    from datetime import date, timedelta
    today = today or date.today()
    text = (text or "").strip()
    if not text:
        return ""
    if "今天" in text:
        return today.isoformat()
    if "明天" in text:
        return (today + timedelta(days=1)).isoformat()
    if "后天" in text:
        return (today + timedelta(days=2)).isoformat()
    m = DATE_RE.search(text)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return date(y, mo, d).isoformat()
        except ValueError:
            return ""
    m = MD_RE.search(text)
    if m:
        mo, d = int(m.group(1)), int(m.group(2))
        for year in (today.year, today.year + 1):
            try:
                cand = date(year, mo, d)
            except ValueError:
                continue
            if cand >= today:
                return cand.isoformat()
        return ""
    return ""


def normalize_time(text: str) -> str:
    """把各种写法归一成 HH:MM; 认不出返回空串(当全天)。"""
    text = (text or "").strip()
    if not text:
        return ""
    m = TIME_RE.search(text)
    if not m:
        return ""
    hour, minute = int(m.group(1)), int(m.group(2) or 0)
    if "下午" in text and hour < 12:
        hour += 12
    if "晚上" in text and hour < 12:
        hour += 12
    if hour > 23 or minute > 59:
        return ""
    return f"{hour:02d}:{minute:02d}"


IMPORT_PROMPT = """下面是我的一份日程/计划文件的内容。请把里面**所有的**安排都用 add_event 记下来。

要求：
- 一条都不要漏，也不要添加文件里根本没有的东西
- 相对日期（下周三、月底、明天）按今天换算成绝对日期
- **不要自己判断某条是不是已经过期/已经办过** —— 看到带日期的活动就记下来，
  过期的部分程序会过滤掉。你自己筛会漏，而漏掉的机会补不回来
- 每条的"事项"写清楚是什么事，别只写"开会"这种看不出来干什么的
- 记完只回一句「共记了 N 条」，不要把清单念一遍

文件内容：
---
{text}
---"""


def collect_by_llm(text: str, ai, pet=None, on_done=None,
                   timeout_hint: int = 120) -> None:
    """让大模型把文本解析成事件 —— **写进临时文件再读出来**, 不碰真实日程。

    这里有个刻意的复用: 模型看到的是和平时**一模一样**的 add_event schema,
    所以参数校验、日期换算、错误兜底全都是同一套代码, 不用为导入再写一遍规则。
    区别只在写入目标 —— 一个临时 Schedule, 读完就删。

    `on_done(events, err_msg, reply)` 在**后台线程**被调用, 调用方自己切回主线程。
    """
    import tempfile

    import pet_schedule
    from ai_tools import Toolbox

    path = os.path.join(tempfile.gettempdir(), f"photopet_import_{os.getpid()}.json")
    if os.path.exists(path):
        os.remove(path)
    scratch = pet_schedule.Schedule(path)

    def done(full: str, err: str):
        events = [{"title": e.title, "date": e.date, "time": e.time,
                   "repeat": e.repeat, "note": e.note} for e in scratch.events]
        try:
            os.remove(path)                     # 临时文件用完就删, 别留垃圾
        except OSError:
            pass
        if on_done:
            on_done(events, err, full)

    ai.chat_stream(IMPORT_PROMPT.format(text=text[:MAX_TEXT_CHARS]),
                   on_done=done, toolbox=Toolbox(scratch, pet))


def normalize_repeat(text: str) -> str:
    text = (text or "").strip()
    if not text:
        return "none"
    table = (("每天", "daily"), ("每日", "daily"), ("daily", "daily"),
             ("每周", "weekly"), ("每星期", "weekly"), ("weekly", "weekly"),
             ("每月", "monthly"), ("monthly", "monthly"),
             ("每年", "yearly"), ("yearly", "yearly"))
    for key, value in table:
        if key in text:
            return value
    return "none"
